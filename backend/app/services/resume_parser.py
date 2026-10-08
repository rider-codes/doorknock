"""Resume file -> text -> structured profile."""
import io

from .. import llm
from ..schemas import ProfileData

SYSTEM = """You read resumes and extract a structured profile.
Rules:
- Extract only what the resume states. Never invent employers, dates, skills or numbers.
- seniority: intern, entry (0-2 years), mid (3-5), senior (6+), staff (leads/architects).
- years_experience: total professional years, counting internships as half. Use 0 if unclear.
- Keep each bullet as written (trim whitespace only)."""


def space_score(text: str) -> float:
    """How wrong the word spacing looks. 0 is ideal. English text is roughly 13-18% spaces with short words;
    PDFs that pack letters tightly come out as 'Builtafinanceappwherepeople...' (under 1% spaces); too-tight reading
    instead shatters words into single letters."""
    words = text.split()
    if not words:
        return 99.0
    space_ratio = text.count(" ") / max(1, len(text))
    singles = sum(1 for w in words if len(w) == 1 and w.isalpha() and w.lower() not in "ai") / len(words)
    return abs(space_ratio - 0.14) * 10 + singles * 4


def best_extraction(extract) -> str:
    """`extract(**settings) -> str` reads the whole document. Try the default first; if its spacing looks off, try tighter
    letter-gap settings and keep whichever reads most like normal text."""
    best_text, best = "", 99.0
    for settings in ({}, {"x_tolerance": 2}, {"x_tolerance": 1.5}, {"x_tolerance": 1}, {"x_tolerance_ratio": 0.1}):
        text = extract(**settings)
        score = space_score(text)
        if score < best:
            best_text, best = text, score
        if settings == {} and score < 0.5:  # the default reading is fine: no need to try others
            break
    return best_text


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        import pdfplumber

        with pdfplumber.open(io.BytesIO(data)) as pdf:
            text = best_extraction(lambda **kw: "\n".join((page.extract_text(**kw) or "") for page in pdf.pages))
    elif name.endswith(".docx"):
        import docx

        doc = docx.Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    else:
        raise ValueError("Upload a PDF or a Word (.docx) file.")
    text = text.strip()
    if len(text) < 80:
        raise ValueError(
            "I could not read any text from that file. If it is a scanned PDF, export a text-based PDF or upload the .docx."
        )
    return text


def parse_profile(raw_text: str) -> ProfileData:
    return llm.structured_model(
        ProfileData,
        purpose="resume_parse",
        system=SYSTEM,
        user="Resume text:\n\n" + raw_text[:24000],
    )
